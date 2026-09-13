from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl


@dataclass
class ExcelRow:
    excel_row: int
    date: object
    product: object
    name: object
    payment_type: object
    phone: object
    result: object = ""


class ExcelReader:
    REQUIRED_COLUMNS = {
        "Дата": "date",
        "Продукт": "product",
        "Имя": "name",
        "Тип Оплаты": "payment_type",
        "Номер": "phone",
    }

    RESULT_COLUMN = "Результат"

    def __init__(self, path: Path) -> None:
        self.path = Path(path)

    # =========================================================
    # HEADER
    # =========================================================

    @staticmethod
    def _normalize_header(value: object) -> str:
        if value is None:
            return ""

        return str(value).strip()

    # =========================================================
    # RESULT COLUMN
    # =========================================================

    def _find_result_column(
        self,
        worksheet,
        headers: dict[str, int],
    ) -> int | None:
        """
        Находит колонку Результат.

        Основной вариант:
            заголовок 'Результат'

        Fallback:
            если заголовка нет, но после обязательных
            колонок существует дополнительная колонка
            с данными, используем её как результат.

        Это нужно для старого Sverka.xlsx, где данные
        результата находятся в F, но заголовок F может
        отсутствовать/быть повреждён.
        """

        # -----------------------------------------------------
        # 1. Нормальный вариант.
        # -----------------------------------------------------

        result_column = headers.get(
            self.RESULT_COLUMN
        )

        if result_column is not None:
            return result_column

        # -----------------------------------------------------
        # 2. Fallback.
        #
        # Обязательные колонки:
        #
        # A = Дата
        # B = Продукт
        # C = Имя
        # D = Тип Оплаты
        # E = Номер
        #
        # Следующая колонка = F.
        # -----------------------------------------------------

        required_columns = [
            headers[column]
            for column in self.REQUIRED_COLUMNS
        ]

        if not required_columns:
            return None

        next_column = max(
            required_columns
        ) + 1

        # Проверяем, существует ли эта колонка.
        if next_column > worksheet.max_column:
            return None

        # Проверяем, есть ли в ней реальные данные.
        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):
            value = worksheet.cell(
                row=row_number,
                column=next_column,
            ).value

            if value is not None and str(value).strip():
                return next_column

        return None

    # =========================================================
    # LOAD
    # =========================================================

    def load(self) -> list[ExcelRow]:
        if not self.path.exists():
            raise FileNotFoundError(
                f"Excel-файл не найден: {self.path}"
            )

        workbook = openpyxl.load_workbook(
            self.path,
            data_only=False,
        )

        try:
            worksheet = workbook.active

            # -------------------------------------------------
            # Headers
            # -------------------------------------------------

            headers: dict[str, int] = {}

            for index, cell in enumerate(
                worksheet[1],
                start=1,
            ):
                header = self._normalize_header(
                    cell.value
                )

                if header:
                    headers[header] = index

            # -------------------------------------------------
            # Required columns
            # -------------------------------------------------

            missing = [
                column
                for column in self.REQUIRED_COLUMNS
                if column not in headers
            ]

            if missing:
                raise ValueError(
                    "В Excel отсутствуют "
                    f"обязательные колонки: {missing}"
                )

            # -------------------------------------------------
            # Result column
            # -------------------------------------------------

            result_column = self._find_result_column(
                worksheet,
                headers,
            )

            rows: list[ExcelRow] = []

            # -------------------------------------------------
            # Read rows
            # -------------------------------------------------

            for row_number in range(
                2,
                worksheet.max_row + 1,
            ):
                result = ""

                if result_column is not None:
                    result = worksheet.cell(
                        row=row_number,
                        column=result_column,
                    ).value

                rows.append(
                    ExcelRow(
                        excel_row=row_number,
                        date=worksheet.cell(
                            row=row_number,
                            column=headers["Дата"],
                        ).value,
                        product=worksheet.cell(
                            row=row_number,
                            column=headers["Продукт"],
                        ).value,
                        name=worksheet.cell(
                            row=row_number,
                            column=headers["Имя"],
                        ).value,
                        payment_type=worksheet.cell(
                            row=row_number,
                            column=headers["Тип Оплаты"],
                        ).value,
                        phone=worksheet.cell(
                            row=row_number,
                            column=headers["Номер"],
                        ).value,
                        result=result,
                    )
                )

            return rows

        finally:
            workbook.close()