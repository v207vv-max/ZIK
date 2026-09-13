from __future__ import annotations

import logging
from pathlib import Path

import openpyxl

logger = logging.getLogger("ZIK.ExcelWriter")


class ExcelWriter:
    RESULT_COLUMN = "Результат"

    def __init__(self, path: Path) -> None:
        self.path = Path(path)

    def _open(self):
        if not self.path.exists():
            raise FileNotFoundError(
                f"Excel-файл не найден: {self.path}"
            )

        return openpyxl.load_workbook(self.path)

    def ensure_result_column(self, worksheet) -> int:
        """
        Находит колонку 'Результат'.
        Если её нет — создаёт справа.
        """

        for cell in worksheet[1]:
            if str(cell.value).strip() == self.RESULT_COLUMN:
                return cell.column

        result_column = worksheet.max_column + 1

        worksheet.cell(
            row=1,
            column=result_column,
            value=self.RESULT_COLUMN,
        )

        logger.info(
            "EXCEL: создана колонка '%s' | column=%s",
            self.RESULT_COLUMN,
            result_column,
        )

        return result_column

    def write_result(
        self,
        excel_row: int,
        result: str,
    ) -> None:
        """
        Записывает результат звонка
        в указанную строку Excel.
        """

        if excel_row < 2:
            raise ValueError(
                f"Некорректная строка Excel: {excel_row}"
            )

        workbook = self._open()

        try:
            worksheet = workbook.active

            result_column = self.ensure_result_column(
                worksheet
            )

            worksheet.cell(
                row=excel_row,
                column=result_column,
                value=result,
            )

            workbook.save(self.path)

            logger.info(
                "EXCEL: записан результат | row=%s | result=%s",
                excel_row,
                result,
            )

        finally:
            workbook.close()